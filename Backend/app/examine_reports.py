#!/usr/bin/env python3
"""
Script to examine the generated Excel reports
"""

import pandas as pd
import os
from openpyxl import load_workbook

def examine_excel_report(filename):
    """Examine an Excel report and show its structure"""
    if not os.path.exists(filename):
        print(f"❌ File {filename} not found!")
        return
    
    print(f"\n📊 EXAMINING REPORT: {filename}")
    print("=" * 80)
    
    try:
        # Load workbook to get sheet names
        workbook = load_workbook(filename)
        sheet_names = workbook.sheetnames
        
        print(f"📄 Total sheets: {len(sheet_names)}")
        print(f"📝 Sheet names: {sheet_names}")
        
        for i, sheet_name in enumerate(sheet_names, 1):
            print(f"\n📋 SHEET {i}: {sheet_name}")
            print("-" * 60)
            
            try:
                # Read the sheet
                df = pd.read_excel(filename, sheet_name=sheet_name)
                
                print(f"📐 Dimensions: {df.shape[0]} rows × {df.shape[1]} columns")
                
                if not df.empty:
                    print(f"📊 Columns: {list(df.columns)}")
                    
                    # Show first few rows
                    print("\n🔍 First 5 rows:")
                    print(df.head().to_string())
                    
                    # Show basic statistics for numeric columns
                    numeric_cols = df.select_dtypes(include=['number']).columns
                    if len(numeric_cols) > 0:
                        print(f"\n📈 Numeric columns statistics:")
                        print(df[numeric_cols].describe().to_string())
                else:
                    print("⚠️ Sheet is empty")
                    
            except Exception as e:
                print(f"❌ Error reading sheet {sheet_name}: {str(e)}")
        
        # Get file size
        file_size = os.path.getsize(filename)
        print(f"\n💾 File size: {file_size:,} bytes ({file_size/1024:.1f} KB)")
        
    except Exception as e:
        print(f"❌ Error examining file: {str(e)}")

def main():
    print("🔍 Excel Report Analysis Tool")
    print("=" * 50)
    
    # Find all recent Excel files
    excel_files = []
    for file in os.listdir('.'):
        if file.endswith('.xlsx') and ('analysis' in file or 'magic' in file):
            excel_files.append(file)
    
    if not excel_files:
        print("❌ No analysis Excel files found!")
        return
    
    # Sort by modification time (newest first)
    excel_files.sort(key=lambda x: os.path.getmtime(x), reverse=True)
    
    print(f"📄 Found {len(excel_files)} analysis reports:")
    for i, file in enumerate(excel_files, 1):
        mtime = os.path.getmtime(file)
        size = os.path.getsize(file)
        print(f"  {i}. {file} ({size:,} bytes)")
    
    # Examine each file
    for file in excel_files:
        examine_excel_report(file)

if __name__ == "__main__":
    main()
