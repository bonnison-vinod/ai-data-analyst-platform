import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.chart.axis import DateAxis
from openpyxl.drawing.image import Image
from openpyxl.worksheet.table import Table, TableStyleInfo
from io import BytesIO
import os
import json
from datetime import datetime
import plotly.graph_objects as go
import plotly.express as px
from plotly.subplots import make_subplots
import base64
import tempfile

class EnhancedExcelDashboard:
    """Enhanced Excel Dashboard Generator with Power BI-like features"""
    
    def __init__(self):
        self.wb = Workbook()
        self.primary_color = "366092"
        self.secondary_color = "D69E2E"
        self.accent_color = "4A90E2"
        self.text_color = "2D3748"
        
        # Define common styles
        self.header_font = Font(name='Calibri', size=14, bold=True, color=self.text_color)
        self.title_font = Font(name='Calibri', size=16, bold=True, color=self.primary_color)
        self.body_font = Font(name='Calibri', size=11, color=self.text_color)
        self.highlight_fill = PatternFill(start_color=self.accent_color, end_color=self.accent_color, fill_type="solid")
        self.header_fill = PatternFill(start_color=self.primary_color, end_color=self.primary_color, fill_type="solid")
        
    def create_dashboard_with_multiple_visuals(self, df, analysis_result, query=None, session_id=None):
        """Create comprehensive dashboard with multiple visualizations"""
        
        # Remove default sheet and create new ones
        self.wb.remove(self.wb.active)
        
        # 1. Executive Summary Dashboard
        self._create_executive_summary(df, analysis_result, query, session_id)
        
        # 2. Data Overview Sheet
        self._create_data_overview(df)
        
        # 3. Detailed Analysis Sheet
        self._create_detailed_analysis(df, analysis_result)
        
        # 4. Interactive Charts Sheet
        self._create_charts_sheet(df, analysis_result)
        
        # 5. Pivot Tables Sheet
        self._create_pivot_tables(df, analysis_result)
        
        # 6. Statistical Summary Sheet
        self._create_statistical_summary(df)
        
        # 7. Insights & Recommendations Sheet
        self._create_insights_sheet(analysis_result, query)
        
        # Save to BytesIO
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output
    
    def _create_executive_summary(self, df, analysis_result, query, session_id):
        """Create executive summary dashboard"""
        ws = self.wb.create_sheet("📊 Executive Summary")
        
        # Header
        ws.merge_cells('A1:H1')
        ws['A1'] = "AI-POWERED DATA ANALYTICS DASHBOARD"
        ws['A1'].font = Font(name='Calibri', size=20, bold=True, color=self.primary_color)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Metadata section
        ws['A3'] = "Analysis Overview"
        ws['A3'].font = self.title_font
        
        metadata = [
            ["Query:", query or "General Analysis"],
            ["Session ID:", session_id or "N/A"],
            ["Generated:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Records Analyzed:", len(df)],
            ["Columns:", len(df.columns)],
            ["Analysis Type:", analysis_result.get('analysis_type', 'Comprehensive')]
        ]
        
        for i, (label, value) in enumerate(metadata):
            ws[f'A{4+i}'] = label
            ws[f'B{4+i}'] = str(value)
            ws[f'A{4+i}'].font = self.header_font
            ws[f'B{4+i}'].font = self.body_font
        
        # Key Metrics Cards
        ws['D3'] = "Key Performance Indicators"
        ws['D3'].font = self.title_font
        
        numeric_cols = df.select_dtypes(include=[np.number]).columns
        if len(numeric_cols) > 0:
            col = numeric_cols[0]
            kpis = [
                ["Total Records", len(df)],
                ["Average " + col.title(), f"{df[col].mean():.2f}"],
                ["Maximum " + col.title(), f"{df[col].max():.2f}"],
                ["Data Quality", f"{(1-df.isnull().sum().sum()/(len(df)*len(df.columns)))*100:.1f}%"]
            ]
            
            for i, (label, value) in enumerate(kpis):
                row = 4 + i
                ws[f'D{row}'] = label
                ws[f'E{row}'] = str(value)
                ws[f'D{row}'].font = self.header_font
                ws[f'E{row}'].font = Font(name='Calibri', size=14, bold=True, color=self.secondary_color)
                ws[f'D{row}'].fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
                ws[f'E{row}'].fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
        
        # Top Insights
        ws['A12'] = "🔍 Key Insights"
        ws['A12'].font = self.title_font
        
        insights = analysis_result.get('insights', [])
        if insights:
            for i, insight in enumerate(insights[:5]):  # Top 5 insights
                ws[f'A{13+i}'] = f"• {insight}"
                ws[f'A{13+i}'].font = self.body_font
                ws[f'A{13+i}'].alignment = Alignment(wrap_text=True)
        
        # Performance Summary Chart (if pivot data available)
        if 'pivot_data' in analysis_result and analysis_result['pivot_data']:
            self._add_summary_chart(ws, analysis_result['pivot_data'])
        
        # Format column widths
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']:
            ws.column_dimensions[col].width = 20
    
    def _create_data_overview(self, df):
        """Create data overview sheet"""
        ws = self.wb.create_sheet("📋 Data Overview")
        
        # Title
        ws['A1'] = "Data Overview & Quality Assessment"
        ws['A1'].font = self.title_font
        
        # Raw data preview
        ws['A3'] = "Data Preview (First 1000 rows)"
        ws['A3'].font = self.header_font
        
        # Add data with formatting
        for r in dataframe_to_rows(df.head(1000), index=False, header=True):
            ws.append(r)
        
        # Format headers
        for col in range(1, len(df.columns) + 1):
            cell = ws.cell(row=4, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Create table
        table_range = f"A4:{chr(64+len(df.columns))}{4+len(df.head(1000))}"
        table = Table(displayName="DataTable", ref=table_range)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                             showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        ws.add_table(table)
        
        # Data quality assessment
        quality_start_row = 4 + len(df.head(1000)) + 3
        ws[f'A{quality_start_row}'] = "Data Quality Assessment"
        ws[f'A{quality_start_row}'].font = self.title_font
        
        # Missing values analysis
        missing_data = []
        for col in df.columns:
            missing_count = df[col].isnull().sum()
            missing_pct = (missing_count / len(df)) * 100
            missing_data.append([col, missing_count, f"{missing_pct:.1f}%"])
        
        quality_headers = ["Column", "Missing Values", "Missing %"]
        ws.append(quality_headers)
        
        for row in missing_data:
            ws.append(row)
    
    def _create_detailed_analysis(self, df, analysis_result):
        """Create detailed analysis sheet"""
        ws = self.wb.create_sheet("📈 Detailed Analysis")
        
        ws['A1'] = "Detailed Analysis Results"
        ws['A1'].font = self.title_font
        
        # Analysis results
        if 'pivot_data' in analysis_result:
            ws['A3'] = "Analysis Results"
            ws['A3'].font = self.header_font
            
            pivot_data = analysis_result['pivot_data']
            if pivot_data:
                pivot_df = pd.DataFrame(pivot_data)
                for r in dataframe_to_rows(pivot_df, index=False, header=True):
                    ws.append(r)
        
        # Summary statistics
        if 'summary_data' in analysis_result:
            summary_start = ws.max_row + 3
            ws[f'A{summary_start}'] = "Summary Statistics"
            ws[f'A{summary_start}'].font = self.header_font
            
            summary_data = analysis_result['summary_data']
            if 'summary_stats' in summary_data:
                for metric, stats in summary_data['summary_stats'].items():
                    ws.append([f"{metric} Statistics"])
                    for stat_name, stat_value in stats.items():
                        ws.append([stat_name.title(), f"{stat_value:,.2f}"])
                    ws.append([])  # Empty row
    
    def _create_charts_sheet(self, df, analysis_result):
        """Create interactive charts sheet"""
        ws = self.wb.create_sheet("📊 Charts & Visuals")
        
        ws['A1'] = "Interactive Charts & Visualizations"
        ws['A1'].font = self.title_font
        
        # Create multiple chart types based on data
        if 'pivot_data' in analysis_result and analysis_result['pivot_data']:
            chart_data = pd.DataFrame(analysis_result['pivot_data'])
            
            # Bar Chart
            if len(chart_data.columns) >= 2:
                self._create_bar_chart(ws, chart_data, start_row=5)
            
            # Line Chart (if time series data)
            if any('date' in col.lower() or 'time' in col.lower() for col in chart_data.columns):
                self._create_line_chart(ws, chart_data, start_row=20)
            
            # Pie Chart (for categorical data)
            if len(chart_data) <= 10:  # Pie charts work best with limited categories
                self._create_pie_chart(ws, chart_data, start_row=35)
    
    def _create_pivot_tables(self, df, analysis_result):
        """Create pivot tables sheet"""
        ws = self.wb.create_sheet("🔄 Pivot Tables")
        
        ws['A1'] = "Pivot Tables & Cross-Tabulations"
        ws['A1'].font = self.title_font
        
        # Create multiple pivot tables
        numeric_cols = df.select_dtypes(include=[np.number]).columns
        categorical_cols = df.select_dtypes(include=['object']).columns
        
        if len(numeric_cols) > 0 and len(categorical_cols) > 0:
            # Basic pivot table
            try:
                pivot = df.pivot_table(
                    values=numeric_cols[0],
                    index=categorical_cols[0],
                    aggfunc='mean',
                    fill_value=0
                ).reset_index()
                
                ws['A3'] = f"Pivot: {categorical_cols[0]} vs {numeric_cols[0]}"
                ws['A3'].font = self.header_font
                
                for r in dataframe_to_rows(pivot, index=False, header=True):
                    ws.append(r)
            except Exception as e:
                ws['A3'] = f"Error creating pivot table: {str(e)}"
        
        # If we have pivot data from analysis, add it
        if 'pivot_data' in analysis_result:
            pivot_start = ws.max_row + 3
            ws[f'A{pivot_start}'] = "Analysis Pivot Data"
            ws[f'A{pivot_start}'].font = self.header_font
            
            pivot_data = analysis_result['pivot_data']
            if pivot_data:
                pivot_df = pd.DataFrame(pivot_data)
                for r in dataframe_to_rows(pivot_df, index=False, header=True):
                    ws.append(r)
    
    def _create_statistical_summary(self, df):
        """Create statistical summary sheet"""
        ws = self.wb.create_sheet("📊 Statistics")
        
        ws['A1'] = "Statistical Summary"
        ws['A1'].font = self.title_font
        
        # Descriptive statistics
        numeric_cols = df.select_dtypes(include=[np.number]).columns
        if len(numeric_cols) > 0:
            desc_stats = df[numeric_cols].describe()
            
            ws['A3'] = "Descriptive Statistics"
            ws['A3'].font = self.header_font
            
            for r in dataframe_to_rows(desc_stats, index=True, header=True):
                ws.append(r)
        
        # Correlation matrix
        if len(numeric_cols) > 1:
            corr_start = ws.max_row + 3
            ws[f'A{corr_start}'] = "Correlation Matrix"
            ws[f'A{corr_start}'].font = self.header_font
            
            corr_matrix = df[numeric_cols].corr()
            for r in dataframe_to_rows(corr_matrix, index=True, header=True):
                ws.append(r)
    
    def _create_insights_sheet(self, analysis_result, query):
        """Create insights and recommendations sheet"""
        ws = self.wb.create_sheet("💡 Insights & Recommendations")
        
        ws['A1'] = "AI-Generated Insights & Recommendations"
        ws['A1'].font = self.title_font
        
        # Query context
        ws['A3'] = "Analysis Query:"
        ws['A3'].font = self.header_font
        ws['A4'] = query or "General data analysis"
        ws['A4'].font = self.body_font
        ws['A4'].alignment = Alignment(wrap_text=True)
        
        # Insights
        insights = analysis_result.get('insights', [])
        if insights:
            ws['A6'] = "Key Insights:"
            ws['A6'].font = self.header_font
            
            for i, insight in enumerate(insights):
                ws[f'A{7+i}'] = f"• {insight}"
                ws[f'A{7+i}'].font = self.body_font
                ws[f'A{7+i}'].alignment = Alignment(wrap_text=True)
        
        # Recommendations (generated based on analysis type)
        recommendations = self._generate_recommendations(analysis_result)
        if recommendations:
            rec_start = ws.max_row + 3
            ws[f'A{rec_start}'] = "Recommendations:"
            ws[f'A{rec_start}'].font = self.header_font
            
            for i, rec in enumerate(recommendations):
                ws[f'A{rec_start+1+i}'] = f"• {rec}"
                ws[f'A{rec_start+1+i}'].font = self.body_font
                ws[f'A{rec_start+1+i}'].alignment = Alignment(wrap_text=True)
        
        # Format column width
        ws.column_dimensions['A'].width = 80
    
    def _add_summary_chart(self, ws, pivot_data):
        """Add summary chart to executive summary"""
        if not pivot_data:
            return
            
        chart_df = pd.DataFrame(pivot_data)
        if len(chart_df.columns) < 2:
            return
        
        # Create chart data starting from column F
        chart_start_row = 12
        chart_start_col = 6  # Column F
        
        # Add chart title
        ws.cell(row=chart_start_row, column=chart_start_col, value="Performance Summary")
        ws.cell(row=chart_start_row, column=chart_start_col).font = self.header_font
        
        # Add chart data
        for r_idx, row in enumerate(dataframe_to_rows(chart_df.head(10), index=False, header=True)):
            for c_idx, value in enumerate(row):
                ws.cell(row=chart_start_row + 2 + r_idx, column=chart_start_col + c_idx, value=value)
        
        # Create Excel chart
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Top Performance"
        chart.y_axis.title = 'Value'
        chart.x_axis.title = 'Category'
        
        # Define data range
        data_range = Reference(ws, min_col=chart_start_col+1, min_row=chart_start_row+2, 
                             max_col=chart_start_col+len(chart_df.columns)-1, max_row=chart_start_row+2+len(chart_df.head(10)))
        cats = Reference(ws, min_col=chart_start_col, min_row=chart_start_row+3, max_row=chart_start_row+2+len(chart_df.head(10)))
        
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(cats)
        
        # Add chart to worksheet
        ws.add_chart(chart, f"H{chart_start_row}")
    
    def _create_bar_chart(self, ws, data, start_row):
        """Create bar chart"""
        if len(data.columns) < 2:
            return
            
        # Add data for chart
        for r_idx, row in enumerate(dataframe_to_rows(data.head(10), index=False, header=True)):
            for c_idx, value in enumerate(row):
                ws.cell(row=start_row + r_idx, column=1 + c_idx, value=value)
        
        # Create chart
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Bar Chart Analysis"
        chart.y_axis.title = 'Value'
        chart.x_axis.title = 'Category'
        
        data_range = Reference(ws, min_col=2, min_row=start_row, max_col=len(data.columns), max_row=start_row+len(data.head(10)))
        cats = Reference(ws, min_col=1, min_row=start_row+1, max_row=start_row+len(data.head(10)))
        
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(cats)
        
        ws.add_chart(chart, f"E{start_row}")
    
    def _create_line_chart(self, ws, data, start_row):
        """Create line chart"""
        if len(data.columns) < 2:
            return
            
        chart = LineChart()
        chart.title = "Trend Analysis"
        chart.style = 13
        chart.y_axis.title = 'Value'
        chart.x_axis.title = 'Time'
        
        # Add data
        for r_idx, row in enumerate(dataframe_to_rows(data.head(20), index=False, header=True)):
            for c_idx, value in enumerate(row):
                ws.cell(row=start_row + r_idx, column=1 + c_idx, value=value)
        
        data_range = Reference(ws, min_col=2, min_row=start_row, max_col=len(data.columns), max_row=start_row+len(data.head(20)))
        cats = Reference(ws, min_col=1, min_row=start_row+1, max_row=start_row+len(data.head(20)))
        
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(cats)
        
        ws.add_chart(chart, f"E{start_row}")
    
    def _create_pie_chart(self, ws, data, start_row):
        """Create pie chart"""
        if len(data.columns) < 2 or len(data) > 10:
            return
            
        chart = PieChart()
        chart.title = "Distribution Analysis"
        
        # Add data
        for r_idx, row in enumerate(dataframe_to_rows(data, index=False, header=True)):
            for c_idx, value in enumerate(row):
                ws.cell(row=start_row + r_idx, column=1 + c_idx, value=value)
        
        data_range = Reference(ws, min_col=2, min_row=start_row+1, max_row=start_row+len(data))
        cats = Reference(ws, min_col=1, min_row=start_row+1, max_row=start_row+len(data))
        
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(cats)
        
        ws.add_chart(chart, f"E{start_row}")
    
    def _generate_recommendations(self, analysis_result):
        """Generate recommendations based on analysis results"""
        recommendations = []
        
        # Basic recommendations based on analysis type
        if 'pivot_data' in analysis_result and analysis_result['pivot_data']:
            recommendations.append("Focus on top-performing categories identified in the analysis")
            recommendations.append("Investigate underperforming areas for improvement opportunities")
            recommendations.append("Consider resource reallocation based on performance insights")
        
        if 'insights' in analysis_result:
            recommendations.append("Implement monitoring for key metrics identified in insights")
            recommendations.append("Set up alerts for performance thresholds")
        
        recommendations.append("Schedule regular analysis updates to track progress")
        recommendations.append("Share insights with relevant stakeholders for action planning")
        
        return recommendations

def create_excel_with_dashboards_and_charts(df, analysis_result, query=None, session_id=None):
    """Main function to create enhanced Excel dashboard"""
    dashboard = EnhancedExcelDashboard()
    return dashboard.create_dashboard_with_multiple_visuals(df, analysis_result, query, session_id)

# Backward compatibility
def create_excel_with_pivots_and_charts(df, analysis_result, query=None, session_id=None):
    """Backward compatibility function"""
    return create_excel_with_dashboards_and_charts(df, analysis_result, query, session_id)
