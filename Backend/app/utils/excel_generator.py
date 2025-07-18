import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from io import BytesIO

def create_excel_with_pivots_and_charts(df, analysis_result, query=None, session_id=None):
    wb = Workbook()
    
    # Add data sheet
    ws_data = wb.active
    ws_data.title = "Data"
    for r in dataframe_to_rows(df, index=False, header=True):
        ws_data.append(r)
        
    # Add a summary sheet with descriptive statistics
    numeric_cols = df.select_dtypes(include=['number']).columns
    if not numeric_cols.empty:
        desc_stats = df[numeric_cols].describe().transpose()
        ws_summary = wb.create_sheet("Data_Summary")
        for r in dataframe_to_rows(desc_stats, index=True, header=True):
            ws_summary.append(r)
            
    # Add analysis insights/story to a sheet
    if analysis_result and isinstance(analysis_result, dict):
        ws_insights = wb.create_sheet("Analysis_Insights")
        ws_insights.append(["AI Generated Insights"])
        ws_insights.append(["Query:", query or "N/A"])
        ws_insights.append(["Session ID:", session_id or "N/A"])
        ws_insights.append(["Generated at:", pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws_insights.append(["Analysis Type:", analysis_result.get('analysis_type', 'N/A')])
        ws_insights.append([""])  # Empty row
        
        insights = analysis_result.get('insights', [])
        if insights:
            ws_insights.append(["Insights:"])
            for insight in insights:
                ws_insights.append([insight])
                
        # Add pivot/analysis data if available
        if 'pivot_data' in analysis_result:
            ws_insights.append([""])  # Empty row
            ws_insights.append(["Analysis Data:"])
            pivot_data = analysis_result['pivot_data']
            if pivot_data:
                # Convert to DataFrame for easier handling
                pivot_df = pd.DataFrame(pivot_data)
                for r in dataframe_to_rows(pivot_df, index=False, header=True):
                    ws_insights.append(r)

    # Save to BytesIO for FastAPI response
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
