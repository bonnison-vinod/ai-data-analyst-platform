#!/usr/bin/env python3
"""
Test script to specifically verify download functionality
"""

import requests
import os
from datetime import datetime

API_BASE = "http://localhost:8000/api"

def test_complete_workflow():
    """Test the complete upload-analyze-download workflow"""
    print("🔬 Testing Complete Upload → Analyze → Download Workflow")
    print("=" * 60)
    
    test_file_path = "Backend/app/uploads/sample_data.csv"
    
    if not os.path.exists(test_file_path):
        print(f"❌ Test file not found: {test_file_path}")
        return
    
    # Step 1: Upload and start analysis
    print("📤 Step 1: Starting analysis...")
    with open(test_file_path, 'rb') as f:
        files = {'file': f}
        data = {
            'question': 'Analyze this sales data and provide comprehensive insights',
            'analysis_type': 'comprehensive',
            'custom_sheets': '[]'
        }
        
        response = requests.post(f"{API_BASE}/generate-enhanced-report/", files=files, data=data)
        
        if response.status_code != 200:
            print(f"❌ Upload failed: {response.status_code}")
            print(response.text)
            return
        
        result = response.json()
        session_id = result['session_id']
        print(f"✅ Analysis started with session ID: {session_id}")
    
    # Step 2: Wait for completion
    print("⏳ Step 2: Waiting for analysis completion...")
    import time
    max_wait = 60  # Max 1 minute
    waited = 0
    
    while waited < max_wait:
        progress_response = requests.get(f"{API_BASE}/report-progress/{session_id}")
        
        if progress_response.status_code == 200:
            progress = progress_response.json()
            percentage = progress.get('percentage', 0)
            status = progress.get('status', 'unknown')
            message = progress.get('message', '')
            
            print(f"📊 Progress: {percentage:.1f}% - {status} - {message}")
            
            if progress.get('completed'):
                if progress.get('error'):
                    print(f"❌ Analysis failed: {progress.get('error')}")
                    return
                else:
                    print("✅ Analysis completed!")
                    break
        
        time.sleep(2)
        waited += 2
    
    if waited >= max_wait:
        print("❌ Analysis timed out")
        return
    
    # Step 3: Test download
    print("⬇️ Step 3: Testing download...")
    download_response = requests.get(f"{API_BASE}/download-report/{session_id}")
    
    print(f"📊 Download Response:")
    print(f"   Status Code: {download_response.status_code}")
    print(f"   Content Type: {download_response.headers.get('content-type', 'Unknown')}")
    print(f"   Content Length: {len(download_response.content)} bytes")
    
    if download_response.status_code == 200:
        # Save the file
        filename = f"test_download_{session_id}.xlsx"
        with open(filename, 'wb') as f:
            f.write(download_response.content)
        
        file_size = os.path.getsize(filename)
        print(f"✅ Download successful!")
        print(f"📄 File saved as: {filename}")
        print(f"📊 File size: {file_size:,} bytes")
        
        # Verify it's a valid Excel file
        try:
            import pandas as pd
            from openpyxl import load_workbook
            
            workbook = load_workbook(filename)
            sheet_names = workbook.sheetnames
            print(f"📋 Excel sheets found: {sheet_names}")
            
            # Try reading first sheet
            df = pd.read_excel(filename, sheet_name=sheet_names[0])
            print(f"📐 First sheet dimensions: {df.shape}")
            
            print("✅ File is a valid Excel workbook!")
            
        except Exception as e:
            print(f"⚠️ File validation warning: {e}")
        
    else:
        print(f"❌ Download failed: {download_response.status_code}")
        print(f"Response: {download_response.text[:200]}")

if __name__ == "__main__":
    test_complete_workflow()
